"""从实际任务/物品配置导出宝箱设计，并同步已有任务、装备资料表。"""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
REPORTS = ROOT / "reports"


def replace_rows(sheet, rows):
    styles = [copy(cell._style) for cell in sheet[2]]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row in rows:
        sheet.append(row)
        for cell, style in zip(sheet[sheet.max_row], styles):
            cell._style = copy(style)
    sheet.auto_filter.ref = sheet.dimensions
    for table in sheet.tables.values():
        table.ref = sheet.dimensions


def new_sheet(workbook, title, headers, rows):
    if title in workbook:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="53634A")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 22
    sheet.row_dimensions[1].height = 32
    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = 36
    if len(headers) == 2:
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 110
        for row_number in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row_number].height = 48
    else:
        sheet.column_dimensions["A"].width = 42
        sheet.column_dimensions["B"].width = 34
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def main():
    missions = yaml.safe_load((ROOT / "data/mission_templates.yaml").read_text())["missions"]
    items = {x["key"]: x for x in yaml.safe_load((ROOT / "data/item_templates.yaml").read_text())["items"]}
    recipes = {
        x["blueprint_key"]: x for x in yaml.safe_load((ROOT / "data/forge_blueprints.yaml").read_text())["recipes"]
    }
    design = {
        x["mission_key"]: x
        for x in json.loads((REPORTS / "mission_reward_design_20261001.json").read_text())["missions"]
    }
    difficulty = {"junior": "初级", "intermediate": "中级", "advanced": "高级"}
    xp_hours = {"experience_fruit": 1, "experience_peach": 8, "experience_pineapple": 16, "experience_watermelon": 24}
    details, budgets, gear_routes, collection_goals = [], [], [], []
    for mission in missions:
        key = mission["key"]
        box_key = next(k for k in mission["drop_table"] if k != "silver")
        box = items[box_key]
        payload = box["effect_payload"]
        entries = [
            ("必得", 1, x["min_quantity"], x["max_quantity"], [{"item_key": x["item_key"], "weight": 1}])
            for x in payload["item_rewards"]
        ]
        entries += [
            (x["label"], x["chance"], x["min_quantity"], x["max_quantity"], x["choices"])
            for x in payload["random_item_groups"]
        ]
        support_entries = []
        for label, chance, minimum, maximum, choices in entries:
            total = sum(c["weight"] for c in choices)
            for choice in choices:
                for material in payload.get("blueprint_material_support", {}).get(choice["item_key"], []):
                    support_entries.append(
                        (
                            "随" + items[choice["item_key"]]["name"] + "附带（同次命中）",
                            chance * choice["weight"] / total,
                            minimum * material["min_quantity"],
                            maximum * material["max_quantity"],
                            [{"item_key": material["item_key"], "weight": 1}],
                        )
                    )
        entries += support_entries
        for label, chance, minimum, maximum, choices in entries:
            total_weight = sum(x["weight"] for x in choices)
            for choice in choices:
                item_key = choice["item_key"]
                item = items[item_key]
                row = len(details) + 2
                recipe = recipes.get(item_key, {})
                result_key = recipe.get("result_item_key", "")
                source = f"data/item_templates.yaml:{box_key}.effect_payload"
                details.append(
                    [
                        key,
                        mission["name"],
                        box_key,
                        box["name"],
                        label,
                        item_key,
                        item["name"],
                        item["rarity"],
                        item["effect_type"],
                        chance,
                        choice["weight"],
                        total_weight,
                        minimum,
                        maximum,
                        f"=J{row}*K{row}/L{row}",
                        f"=(M{row}+N{row})/2",
                        f"=O{row}*P{row}",
                        xp_hours.get(item_key, 0),
                        f"=Q{row}*R{row}",
                        recipe.get("required_forging"),
                        items[result_key]["name"] if result_key else "",
                        source,
                    ]
                )
                probability = chance * choice["weight"] / total_weight
                collection_payload = item.get("effect_payload", {})
                if collection_payload.get("collection_cost"):
                    target_key = collection_payload["result_item_key"]
                    threshold = collection_payload["collection_cost"]
                    target_probability = sum(
                        g["chance"] * c["weight"] / sum(candidate["weight"] for candidate in g["choices"])
                        for g in payload["random_item_groups"]
                        for c in g["choices"]
                        if c["item_key"] == target_key
                    )
                    goal_row = len(collection_goals) + 2
                    collection_goals.append(
                        [
                            key,
                            mission["name"],
                            item["name"],
                            items[target_key]["name"],
                            threshold,
                            1,
                            mission.get("daily_limit", 3),
                            target_probability,
                            f"=(1-(1-H{goal_row})^E{goal_row})/H{goal_row}",
                            f"=(1-H{goal_row})^E{goal_row}",
                            "胜利并开箱累积；仓库使用兑换；不含战败/缺勤；随机掉落不清空碎片",
                        ]
                    )
                    gear_routes.append(
                        [
                            target_key,
                            items[target_key]["name"],
                            "碎片累计兑换（任务宝箱）",
                            f"{mission['name']}胜利开箱必得{item['name']}×1；累计{threshold}个后仓库使用兑换1件，碎片不可交易",
                            source,
                        ]
                    )
                condition = f"任务：{mission['name']}（{key}）；胜利得{box['name']}×1，在仓库开启；组概率{chance:.6%}，权重{choice['weight']}/{total_weight}，单项概率{probability:.6%}，数量{minimum}—{maximum}"
                if item["effect_type"].startswith("equip_"):
                    gear_routes.append([item_key, item["name"], "任务宝箱装备", condition, source])
                elif recipe:
                    materials = "、".join(
                        f"{items.get(k, {}).get('name', k)}×{v}" for k, v in recipe.get("costs", {}).items()
                    )
                    gear_routes.append(
                        [
                            result_key,
                            items[result_key]["name"],
                            "图纸合成（任务宝箱）",
                            condition
                            + f"；获得{item['name']}，锻造{recipe['required_forging']}级，合成产出{recipe.get('quantity_out', 1)}件；材料：{materials}",
                            source + "；data/forge_blueprints.yaml",
                        ]
                    )
        reference = design[key]
        row = len(budgets) + 2
        budgets.append(
            [
                key,
                mission["name"],
                difficulty[mission["difficulty"]],
                box["name"],
                payload["reward_theme"],
                mission["drop_table"].get("silver", 0),
                ",".join(map(str, mission.get("available_weekdays", []))) or "每日",
                mission.get("daily_limit", 3),
                reference["fixed_experience_fruit"],
                reference["old_configured_xp_hours"],
                reference["enemy_salvage_reference_hours"],
                f"=SUMIF('宝箱奖励明细'!A$2:A${len(details)+1},A{row},'宝箱奖励明细'!S$2:S${len(details)+1})",
                f"=L{row}-J{row}",
                json.dumps(mission.get("entry_cost", {}), ensure_ascii=False),
                "箱不可交易；器械仅图纸；仓容5",
            ]
        )
    # SUMIF统一使用完整范围，避免新增任务后引用变化。
    for row_number, row in enumerate(budgets, 2):
        row[11] = f"=SUMIF('宝箱奖励明细'!A$2:A${len(details)+1},A{row_number},'宝箱奖励明细'!S$2:S${len(details)+1})"
    headers = [
        "任务key",
        "任务名称",
        "宝箱key",
        "宝箱名称",
        "奖励组",
        "物品key",
        "物品名称",
        "品质",
        "类型",
        "组命中概率",
        "池内权重",
        "权重合计",
        "最少数量",
        "最多数量",
        "单项实际概率",
        "平均数量",
        "每箱期望数量",
        "单件经验(小时)",
        "每箱经验期望(小时)",
        "所需锻造等级",
        "图纸产物",
        "配置来源",
    ]
    budget_headers = [
        "任务key",
        "任务名称",
        "难度",
        "宝箱名称",
        "养成定位",
        "直接银两",
        "开放星期",
        "每日基础次数",
        "固定经验果",
        "原配置经验期望(小时)",
        "敌军回收参考(小时)",
        "新箱经验期望(小时)",
        "较原配置增加(小时)",
        "入场消耗",
        "规则",
    ]
    rules = [
        ["范围", "32个人任务；3帮会任务维持公共红宝石奖励。"],
        ["发奖", "玩家胜利且非撤退，在既有结算时机获得原银两+对应宝箱×1；原无银两的分桃、地窖二层仍只奖励宝箱。"],
        ["开箱", "仓库点击使用；保底经验+多个独立奖励组，命中某组只按权重选一项；组间可同时命中。"],
        [
            "经验预算",
            "固定经验果取敌方门客全败满血+敌方护院80%阵亡参考向下取整、至少1；保留原配置经验概率和数量。不计玩家损失，不等同旧每场实际收益。",
        ],
        ["经验差值", "新箱减原配置的差值主要是战斗回收归入固定保底；不是凭空新增同等经验。实际新旧差异取决于旧战损。"],
        [
            "器械",
            "只得图纸，命中时附带配方每种材料的一半向下取整；剩余材料与锻造等级仍需满足。半材料与图纸同次命中，不是独立抽取。",
        ],
        ["活动", "月饼、蜡笔固定1及55%×1、25%×2、8%×4、1.5%×10独立奖励组保留；单次期望2.52。月饼兼作兑换与开箱道具。"],
        [
            "任务链",
            "汜水战报、虎牢战旗、阵旗碎片、玉虚残印在对应箱各必得1，开箱取得凭证后用于下一关。宝箱不可交易。",
        ],
        [
            "兼容",
            "已有非空旧掉落战报仍按旧奖励与回收结算；空掉落防守战报沿用领取时当前模板。旧武林箱保留原池/价格/交易/仓容，器械改图纸；新版高级华山用独立箱。",
        ],
        [
            "仓库",
            "箱占5；普通任务/开箱奖励沿用仓满仍入库的既有策略，不吞奖励。缺任一奖励候选模板则保留整个宝箱等待同步。",
        ],
        [
            "概率",
            "概率不含胜率。EV=组概率×权重/总权重×平均数量。华山85%紫/15%橙互斥保证一件；官渡15%橙与25%紫独立；配套材料随图纸命中。",
        ],
        [
            "累计兑换",
            "6个高阶任务各有独立碎片，每箱固定1；诛仙30个，其余40个兑换对应指定橙装。不承诺自然日，随机先出仍保留累计进度。",
        ],
        ["图像", "32箱统一复用mission_chest_v1.png，国风木铜宝箱，很浅的浅黄色背景。"],
        ["来源", "工作区YAML配置与结算、开箱源码，2026-10-01；未读取或修改线上数据库。"],
    ]
    catalog = openpyxl.load_workbook(REPORTS / "mission_catalog.xlsx")
    direct_sheet = catalog["任务掉落明细"]
    guild_rows = [list(row) for row in list(direct_sheet.values)[1:] if row[25] == "帮会任务"]
    direct = []
    summaries = {row[1]: list(row) for row in list(catalog["任务列表"].values)[1:]}
    for mission in missions:
        old = summaries[mission["key"]]
        for item_key, quantity in mission["drop_table"].items():
            row = len(direct) + 2
            item = items.get(item_key, {})
            direct.append(
                [
                    len(direct) + 1,
                    *old[1:11],
                    "固定掉落",
                    item_key,
                    item_key,
                    item.get("name", "银两"),
                    item.get("effect_type", "资源"),
                    item.get("rarity", "资源"),
                    1,
                    quantity,
                    None,
                    "否",
                    str(quantity),
                    None,
                    f"=R{row}",
                    f"=X{row}*S{row}",
                    "个人任务",
                    None,
                    "挑战者庄园",
                    "玩家胜利、非撤退，任务返程/防守结算；宝箱内容另见宝箱奖励明细",
                    f"data/mission_templates.yaml:{mission['key']}.drop_table.{item_key}",
                ]
            )
    for row in guild_rows:
        number = len(direct) + 2
        row[0] = number - 1
        row[23] = f"=R{number}"
        row[24] = f"=X{number}*S{number}"
        direct.append(row)
    replace_rows(direct_sheet, direct)
    end = len(direct) + 1
    for row in catalog["任务列表"].iter_rows(min_row=2):
        for column in (14, 15, 16):
            row[column - 1].value = row[column - 1].value.replace("$290", f"${end}")
        key = row[1].value
        if key in design:
            mission = next(x for x in missions if x["key"] == key)
            row[23].value = "；".join(
                (items.get(k, {}).get("name", "银两")) + f"×{v}" for k, v in mission["drop_table"].items()
            )
            row[24].value = "开箱奖励及概率见宝箱奖励明细"
            row[28].value = f"data/mission_templates.yaml:{key}"
            row[31].value = json.dumps(mission, ensure_ascii=False, separators=(",", ":"))
    reverse = []
    for idx, row in enumerate(direct, 2):
        reverse.append(
            [
                row[13],
                row[14],
                row[15],
                row[16],
                row[1],
                row[2],
                row[11],
                f"='任务掉落明细'!X{idx}",
                row[18],
                row[12],
                row[25],
                row[27],
                f"='任务掉落明细'!Y{idx}",
                row[29],
            ]
        )
    replace_rows(catalog["掉落物反查"], reverse)
    replace_rows(catalog["页面展示配置"], [])
    replace_rows(
        catalog["动态奖励规则"],
        [list(row) for row in list(catalog["动态奖励规则"].values)[1:] if row[0] != "32新版个人任务"],
    )
    for row in catalog["动态奖励规则"].iter_rows(min_row=2, max_row=4):
        row[0].value = "改版前已有散装奖励战报（不含任务宝箱）"
    catalog["动态奖励规则"].append(
        [
            "32新版个人任务",
            "经验与装备统一收入宝箱",
            "对应mission_chest_*",
            "宝箱内容",
            "按每箱配置发奖；不再额外散发战损经验/装备",
            "见宝箱养成预算",
            "胜利非撤退得箱；玩家仓库开启",
            "gameplay/services/missions_impl/finalization_helpers.py；gameplay/services/inventory/use.py",
        ]
    )
    for row in catalog["统计"].iter_rows(min_row=2):
        if isinstance(row[1].value, str):
            row[1].value = row[1].value.replace("B290", f"B{end}")
        if row[0].value == "生成时间":
            row[1].value = "2026-10-01"
    for row in catalog["字段说明"].iter_rows(min_row=2):
        if row[0].value == "奖励掉落范围":
            row[1].value = (
                "新版个人任务直接发银两与宝箱；内部奖励见新增宝箱奖励明细。旧战报回收见动态奖励规则。帮会公共红宝石机制保留。"
            )
    new_sheet(catalog, "宝箱奖励明细", headers, details)
    new_sheet(catalog, "宝箱养成预算", budget_headers, budgets)
    new_sheet(catalog, "宝箱规则说明", ["项目", "说明"], rules)
    goal_headers = [
        "任务key",
        "任务名称",
        "累积道具",
        "兑换装备",
        "兑换需要胜利开箱数",
        "每箱碎片",
        "每日基础次数",
        "装备随机概率",
        "首次获得期望胜利数（含兑换）",
        "达到兑换数仍未随机出货概率",
        "假设与限制",
    ]
    new_sheet(catalog, "长期目标与兑换", goal_headers, collection_goals)
    catalog.save(REPORTS / "mission_catalog.xlsx")
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet = new_sheet(workbook, "宝箱养成预算", budget_headers, budgets)
    for col in (6, 9, 10, 11, 12, 13):
        for cell in list(sheet.columns)[col - 1][1:]:
            cell.number_format = "#,##0.00"
    sheet = new_sheet(workbook, "宝箱奖励明细", headers, details)
    for col in (10, 15):
        for cell in list(sheet.columns)[col - 1][1:]:
            cell.number_format = "0.0000%"
    for col in (17, 19):
        for cell in list(sheet.columns)[col - 1][1:]:
            cell.number_format = "0.000000"
    new_sheet(workbook, "设计与开启规则", ["项目", "说明"], rules)
    goals_sheet = new_sheet(workbook, "长期目标与兑换", goal_headers, collection_goals)
    for col in (8, 10):
        for cell in list(goals_sheet.columns)[col - 1][1:]:
            cell.number_format = "0.00%"
    for cell in list(goals_sheet.columns)[8][1:]:
        cell.number_format = "0.00"
    new_sheet(
        workbook,
        "帮会任务保留",
        ["任务key", "任务名称", "公共红宝石奖励", "接收方"],
        [
            [x["key"], x["name"], x["ruby_reward"], "帮会仓库"]
            for x in yaml.safe_load((ROOT / "data/guild_mission_templates.yaml").read_text())["missions"]
        ],
    )
    workbook.save(REPORTS / "mission_reward_chests_20261001.xlsx")
    equipment = openpyxl.load_workbook(REPORTS / "equipment_catalog.xlsx")
    old_routes = [
        list(row)
        for row in list(equipment["获取途径"].values)[1:]
        if "data/mission_templates.yaml" not in str(row[4])
        and row[2] not in {"任务宝箱装备", "图纸合成（任务宝箱）", "碎片累计兑换（任务宝箱）"}
        and not any(
            key in str(row[3]) for key in ("huashan_wulin_chest", "saiwai_equipment_box", "fengshen_equipment_box")
        )
    ]
    replace_rows(equipment["获取途径"], old_routes + gear_routes)
    available = {row[0] for row in old_routes + gear_routes}
    replace_rows(
        equipment["获取途径待核实"],
        [list(row) for row in list(equipment["获取途径待核实"].values)[1:] if row[0] not in available],
    )
    methods = {}
    for row in old_routes + gear_routes:
        methods.setdefault(row[0], set()).add(row[2])
    for row in equipment["装备总表"].iter_rows(min_row=2):
        row[30].value = "；".join(sorted(methods.get(row[1].value, set()))) or "当前主要配置未发现明确途径"
    for row in equipment["更新说明"].iter_rows(min_row=2):
        if row[0].value == "更新时间":
            row[1].value = "2026-10-01"
        if row[0].value == "途径待核实":
            row[1].value = f"{equipment['获取途径待核实'].max_row-1}件；个人任务途径已同步为宝箱与图纸合成。"
    equipment.save(REPORTS / "equipment_catalog.xlsx")
    print(
        f"exported {len(missions)} missions, {len(details)} chest reward rows, {len(gear_routes)} gear acquisition routes"
    )


if __name__ == "__main__":
    main()
